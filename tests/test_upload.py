"""Uploaded resumes (PDF, DOCX, TXT) turn into text; bad files get a clear error."""
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

import app as app_module
import resume_text
from resume_text import ExtractError, extract_text

LINES = ["Aarav Mehta", "Backend Developer, Finly", "Built payment reconciliation APIs in FastAPI"]


def make_pdf(lines=LINES) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i, line in enumerate(lines):
        c.drawString(72, 750 - 20 * i, line)
    c.save()
    return buf.getvalue()


def make_blank_pdf() -> bytes:          # what a scanned resume looks like to a text extractor
    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    doc = Document()
    for line in LINES:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, PostgreSQL, Redis"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf():
    text = extract_text("resume.pdf", make_pdf())
    assert "Aarav Mehta" in text and "FastAPI" in text


def test_docx_including_tables():
    text = extract_text("Resume.DOCX", make_docx())
    assert "Backend Developer, Finly" in text
    assert "Skills | Python, PostgreSQL, Redis" in text


def test_txt():
    text = extract_text("cv.txt", "\n".join(LINES).encode())
    assert text.splitlines()[0] == "Aarav Mehta"


def test_scanned_pdf_explains_what_to_do():
    with pytest.raises(ExtractError, match="scanned"):
        extract_text("scan.pdf", make_blank_pdf())


def test_corrupt_file():
    with pytest.raises(ExtractError, match="Couldn't read"):
        extract_text("resume.pdf", b"%PDF-1.4 this is not really a pdf")


def test_wrong_type():
    with pytest.raises(ExtractError) as e:
        extract_text("photo.png", b"\x89PNG....")
    assert e.value.status == 415


def test_too_big():
    with pytest.raises(ExtractError) as e:
        extract_text("cv.txt", b"x" * (resume_text.MAX_BYTES + 1))
    assert e.value.status == 413


def test_empty_file():
    with pytest.raises(ExtractError, match="empty"):
        extract_text("cv.txt", b"")


def test_clean_tidies_whitespace():
    assert resume_text.clean("a  \r\n\r\n\r\n\r\nb\x00") == "a\n\nb"


@pytest.fixture
def client():
    return TestClient(app_module.app)


def test_upload_endpoint(client):
    r = client.post("/extract", files={"file": ("resume.pdf", make_pdf(), "application/pdf")})
    assert r.status_code == 200
    assert "Aarav Mehta" in r.json()["text"]


def test_upload_endpoint_rejects_images(client):
    r = client.post("/extract", files={"file": ("photo.png", b"\x89PNG", "image/png")})
    assert r.status_code == 415
    assert "PDF, DOCX or TXT" in r.json()["detail"]
