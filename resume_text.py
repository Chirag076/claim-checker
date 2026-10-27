"""
resume_text.py - turns an uploaded resume (PDF, DOCX or TXT) into plain text.

The text goes back into the resume box on the page, so the candidate can see
exactly what the checker will compare against and fix anything the file
conversion got wrong before running a check.
"""
import io
import re
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024
TYPES = (".pdf", ".docx", ".txt", ".md")
MIN_CHARS = 30      # less than this from a PDF almost always means a scanned image


class ExtractError(ValueError):
    def __init__(self, message, status=422):
        super().__init__(message)
        self.status = status


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ExtractError("This PDF is password protected. Remove the password or paste the text.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:                       # many resume templates are built from tables
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text.strip() and cell.text not in cells:    # merged cells repeat their text
                    cells.append(cell.text)
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in TYPES:
        raise ExtractError("Upload a PDF, DOCX or TXT file.", status=415)
    if len(data) > MAX_BYTES:
        raise ExtractError("Keep the file under 5 MB.", status=413)
    if not data:
        raise ExtractError("The file is empty.")
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except ExtractError:
        raise
    except Exception:
        raise ExtractError("Couldn't read this file. Try another format or paste the text.")
    text = clean(text)
    if len(text) < MIN_CHARS:
        raise ExtractError("No text found in this file. If it's a scanned PDF, paste the text instead.")
    return text
